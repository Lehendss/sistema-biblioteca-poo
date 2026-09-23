"""Genera la memoria Word con tablas nativas y capturas de Flet."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DESTINO = ROOT / "entregas/semanas5y6"
documento = Document()
seccion = documento.sections[0]
seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
seccion.top_margin = seccion.bottom_margin = Cm(2)
seccion.left_margin = seccion.right_margin = Cm(2)
normal = documento.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12
for nombre in ["Title", "Heading 1", "Heading 2"]:
    documento.styles[nombre].font.name = "Arial"
    documento.styles[nombre].font.color.rgb = RGBColor.from_string("17365D")
idioma = OxmlElement("w:lang")
idioma.set(qn("w:val"), "es-EC")
normal.element.get_or_add_rPr().append(idioma)
pie = seccion.footer.paragraphs[0]
pie.text = "Edisson Carchi · Programación Orientada a Objetos · "
campo = OxmlElement("w:fldSimple")
campo.set(qn("w:instr"), "PAGE")
pie._p.append(campo)


def p(texto):
    documento.add_paragraph(texto)


def tabla(cabecera, filas):
    t = documento.add_table(rows=1, cols=len(cabecera))
    t.style = "Light Shading Accent 1"
    for celda, texto in zip(t.rows[0].cells, cabecera):
        celda.text = texto
    for fila in filas:
        for celda, texto in zip(t.add_row().cells, fila):
            celda.text = texto
    for fila in t.rows:
        no_cortar = OxmlElement("w:cantSplit")
        fila._tr.get_or_add_trPr().append(no_cortar)


documento.add_heading("Semanas 5 y 6", 0)
documento.add_heading("Colecciones, genéricos, interfaz gráfica y manejo de eventos", 1)
tabla(["Dato", "Información"], [
    ["Estudiante", "Edisson Carchi"], ["Matrícula", "2025372938"],
    ["Asignatura", "Programación Orientada a Objetos"],
    ["Tecnología", "Python 3.12 · Flet 0.28.3"],
    ["Repositorio", "https://github.com/Lehendss/sistema-biblioteca-poo"]])
documento.add_heading("1. Solución e integración", 1)
p("Se implementó un catálogo gráfico de productos para la biblioteca. Los productos son libros y revistas: "
  "Producto contiene un objeto Libro o Revista, reutilizando las clases derivadas de Material de las semanas anteriores. "
  "Se incorporan precio y stock, y se mantienen atributos privados y acceso controlado mediante propiedades.")
p("CatalogoProductos administra las operaciones CRUD. La interfaz Flet utiliza esa misma instancia para modificar "
  "y consultar los datos. Cada sesión inicia con dos productos de demostración. El almacenamiento es temporal en memoria; "
  "no se incorpora persistencia porque la actividad se centra en las colecciones y los eventos.")
documento.add_heading("2. Colecciones y tipos", 1)
tabla(["Colección", "Aplicación"], [
    ["dict[str, Producto]", "Índice de productos por código; permite buscar, actualizar y eliminar."],
    ["set[str]", "Códigos únicos normalizados. Rechaza duplicados y libera el código al eliminar."],
    ["list[str]", "Historial ordenado de operaciones exitosas. listar() devuelve una lista filtrada de productos."]])
p("En Python se usan anotaciones Dict, Set y List para expresar los tipos esperados. No son validaciones automáticas "
  "en ejecución. La exigencia de genéricos explícitos de la consigna corresponde a Java. "
  "Las consultas devuelven copias para evitar que un consumidor altere las colecciones internas.")
documento.add_page_break()
documento.add_heading("3. Interfaz y manejo de eventos", 1)
tabla(["Control / evento", "Comportamiento"], [
    ["Agregar / on_click", "Valida el formulario, construye el producto y lo incorpora al catálogo."],
    ["Buscador / on_change y on_submit", "Filtra por código o título; vacío muestra todos los productos."],
    ["Editar / on_click", "Consulta un producto y carga sus datos en el formulario."],
    ["Actualizar / on_click", "Reemplaza los datos del producto seleccionado, conservando su código."],
    ["Eliminar / on_click", "Elimina el producto y su código del conjunto de identificadores."],
    ["Tipo / on_change", "Muestra autor e ISBN para libros o edición para revistas."],
    ["Nuevo / limpiar", "Reinicia el formulario y habilita una nueva creación."]])
p("Después de cada operación se actualizan la tabla, el contador, el historial y un mensaje de éxito o error. "
  "Los campos informan sus cambios mediante eventos; no se modifica directamente el diccionario desde la interfaz.")
documento.add_heading("4. Validación y pruebas", 1)
p("Se exige código y título, autor e ISBN para libros, edición entera positiva para revistas, precio positivo finito "
  "con hasta dos decimales y stock entero no negativo. El precio se maneja con Decimal. La identidad se normaliza "
  "quitando espacios externos y convirtiendo a mayúsculas; L001 y l001 representan el mismo código.")
p("Se ejecutaron 14 pruebas unitarias, incluidas las de las actividades anteriores. Las pruebas nuevas comprueban "
  "CRUD, duplicados sin cambios parciales, códigos inexistentes, protección mediante copias, validaciones y conservación de identidad. "
  "Además se ejecutaron interacciones reales mediante Playwright sobre Flet web: alta, búsqueda, actualización, eliminación, "
  "rechazo de duplicados y rechazo de stock negativo. Las figuras siguientes muestran esas ejecuciones.")
documento.add_heading("5. Ejecución", 1)
for comando in ["python3 -m venv .venv", "source .venv/bin/activate", "python -m pip install -r requirements.txt",
                "python -m src.gui", "python -m src.gui --web --port 8550", "python -m unittest discover -s tests -v"]:
    parrafo = documento.add_paragraph(comando)
    parrafo.runs[0].font.name = "Courier New"
    parrafo.runs[0].font.size = Pt(9)
p("Ejecutar desde la raíz del repositorio. En Windows se activa el entorno con .venv\\Scripts\\activate. "
  "La versión web abre http://127.0.0.1:8550. El README contiene las instrucciones completas.")

figuras = [
    ("01_catalogo.png", "Catálogo inicial", "Se muestran los productos L001 y R001, con precio y stock, y los controles CRUD."),
    ("02_agregar.png", "Crear producto", "Se agrega L002 con precio $25.50 y stock 5. El catálogo pasa de dos a tres productos."),
    ("03_buscar.png", "Buscar y listar", "El filtro L002 muestra un resultado sin eliminar los demás productos del catálogo."),
    ("04_actualizar.png", "Actualizar producto", "L002 conserva su código; el precio cambia a $30.00 y el stock a 9."),
    ("05_eliminar.png", "Eliminar producto", "Se elimina L002. Al limpiar el filtro se muestran de nuevo los dos productos iniciales."),
    ("06_duplicado.png", "Evitar duplicados", "El intento de agregar l001 con espacios externos se rechaza por coincidir con L001."),
    ("07_validacion.png", "Validar entradas", "Se rechaza un stock negativo y se mantiene el catálogo sin incorporar el producto inválido.")]
for numero, (archivo, titulo, descripcion) in enumerate(figuras, 1):
    documento.add_page_break()
    documento.add_heading(f"Figura {numero}. {titulo}", 1)
    p(descripcion)
    documento.add_picture(str(DESTINO / "capturas" / archivo), width=Cm(17))
    p("Fuente: ejecución real de la aplicación Flet en Chromium; captura automatizada mediante Playwright.")

documento.add_page_break()
documento.add_heading("6. Conclusión", 1)
p("El proyecto integra las colecciones con una interfaz gráfica funcional. El diccionario facilita la localización "
  "por código, el conjunto impide duplicados y la lista conserva la secuencia de operaciones. Los eventos enlazan "
  "las acciones del usuario con el catálogo sin mezclar el almacenamiento con la presentación. La reutilización de "
  "Libro y Revista mantiene la continuidad del modelo orientado a objetos. Las validaciones y las copias reducen "
  "los estados inconsistentes, mientras que las pruebas comprueban las operaciones y los errores más relevantes. "
  "Una ampliación futura podría incorporar persistencia para conservar el catálogo entre sesiones.")
documento.save(DESTINO / "Carchi_Edisson_Semanas5y6.docx")
print("Memoria Word generada con siete capturas incrustadas.")
